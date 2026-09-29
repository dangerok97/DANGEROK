"""Invito a presentarsi (Vol. IV, All. 14) in Word, identico al modello del Reparto.

Si parte dal file modello (`wordtemplates/invito.docx`, ricavato dall'invito del Reparto): A4, margini, Arial, logo della
Repubblica, intestazione della Compagnia di Tarquinia, numerazione in testata, tabella «OGGETTO», elenchi puntati e
relazione di notifica restano quelli dell'originale. Il programma sostituisce SOLO i dati del caso, mantenendo la
formattazione dei paragrafi. Le due varianti (verifica / controllo) condividono lo stesso impianto e cambiano le formule.
"""
from __future__ import annotations

import copy
import io
import pathlib
import re

import docx
from docx.enum.text import WD_COLOR_INDEX

MODELLO = pathlib.Path(__file__).parent / "wordtemplates" / "invito.docx"

# posizioni dei paragrafi nel modello (contate come nell'originale: solo paragrafi, tabella esclusa)
I_DITTA, I_CF, I_PIVA, I_AL, I_VIA = 7, 8, 9, 11, 12
I_FINE, I_INTRO_DOC, I_DOC_PRIMO, I_DOC_ULTIMO, I_PERIODI, I_CONTATTI = 14, 15, 16, 20, 21, 22
I_RAGIONI, I_OPERAZIONI, I_COMANDANTE, I_NOME_COMANDANTE = 25, 28, 45, 46

OGGETTO = {
    "verifica": "Invito a presentarsi. Avvio di una verifica fiscale ai fini delle imposte sui redditi e/o dell’I.V.A. "
                "e/o degli altri tributi nei confronti di.",
    "controllo": "Invito a presentarsi. Avvio di un controllo fiscale ai fini di P.T. nei confronti di:",
}
NORME_VERIFICA = ("ai sensi dell’art. 32, comma 1, n. 2) e 3) del D.P.R. 29 settembre 1973, n. 600 e/o dell’art. 51, "
                  "comma 2, n. 2) e 3), del D.P.R. 26 ottobre 1972, n. 633 e/o dell’art. 2 del D.Lgs 68/2001")
NORME_CONTROLLO = ("ai fini di P.T., ai sensi e per gli effetti degli artt. 52 e 63 del D.P.R. 26 ottobre 1972, n. 633, "
                   "33 del D.P.R. 29 settembre 1973, n. 600, 2 del D.Lgs 68/2001, nonché della L. n. 4/1929")
RAGIONI_VERIFICA = ("1. le ragioni giustificative dell’invito consistono nel fatto che questo Reparto deve eseguire una "
                    "verifica delle imposte sui redditi e/o dell'I.V.A. e/o degli altri tributi ai sensi dell’art. 32, "
                    "comma 1, n. 2) e 3) del D.P.R. 29 settembre 1973, n. 600 e/o dell’art. 51, comma 2, n. 2) e 3), del "
                    "D.P.R. 26 ottobre 1972, n. 633 in quanto ")
RAGIONI_CONTROLLO = ("1. le ragioni giustificative dell’invito consistono nel fatto che questo Reparto deve eseguire un "
                     "controllo ")
INTRO_RAGIONI = {"verifica": "… deve eseguire una verifica … ai sensi dell’art. 32 … in quanto",
                 "controllo": "… deve eseguire un controllo"}

ORA_VUOTA, DATA_VUOTA = "...", "…………"


def _anni(periodi: list[str]) -> tuple[str, str]:
    """('dal 2019 al 2023', '2019, 2020, 2021, 2022, 2023'): il 'dal … al …' solo se sono anni consecutivi."""
    elenco = ", ".join(periodi)
    anni = [p for p in periodi if re.fullmatch(r"\d{4}", p)]
    if len(anni) == len(periodi) and len(anni) >= 2:
        n = [int(a) for a in anni]
        if n == list(range(n[0], n[-1] + 1)):
            return f"dal {n[0]} al {n[-1]}", elenco
    return elenco, elenco


def _lista(v) -> list[str]:
    if isinstance(v, str):
        v = v.splitlines()
    return [x.strip() for x in (v or []) if str(x).strip()]


def _rpr_modello(p, italic: bool | None = None):
    """Copia le proprieta' del primo carattere del paragrafo (per mantenere Arial, dimensione, ecc.)."""
    for r in p.runs:
        if italic is None or bool(r.italic) == italic:
            rpr = r._r.find(docx.oxml.ns.qn("w:rPr"))
            return copy.deepcopy(rpr) if rpr is not None else None
    for r in p.runs:
        rpr = r._r.find(docx.oxml.ns.qn("w:rPr"))
        return copy.deepcopy(rpr) if rpr is not None else None
    return None


def riscrivi(p, pezzi):
    """Sostituisce il testo del paragrafo mantenendone la formattazione. pezzi: [(testo, stile)], stile '', 'i', 'dc'."""
    base = _rpr_modello(p, italic=False)
    corsivo = _rpr_modello(p, italic=True)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for testo, stile in pezzi:
        run = p.add_run(testo)
        modello = corsivo if (stile == "i" and corsivo is not None) else base
        if modello is not None:
            old = run._r.find(docx.oxml.ns.qn("w:rPr"))
            if old is not None:
                run._r.remove(old)
            run._r.insert(0, copy.deepcopy(modello))
        if stile == "i" and corsivo is None:
            run.italic = True
        if stile == "dc":
            run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE


def _o_da_compilare(valore: str, cosa: str):
    v = (valore or "").strip()
    return (v, "") if v else (f"[DA COMPILARE: {cosa}]", "dc")


def crea_invito(dati: dict, tipo: str, reparto: dict | None = None, modello: pathlib.Path | str | None = None) -> bytes:
    """`tipo`: 'verifica' o 'controllo'. `reparto`: comandante, in_sv, referenti, telefono."""
    assert tipo in ("verifica", "controllo")
    rep = {"comandante": "", "in_sv": True, "referenti": [], "telefono": "0766/856028", **(reparto or {})}
    d = docx.Document(str(modello or MODELLO))
    P = list(d.paragraphs)                                  # riferimenti fissati prima di qualsiasi modifica
    periodi = _lista(dati.get("periodi") if not isinstance(dati.get("periodi"), str) else re.split(r"[;,\n]", dati["periodi"]))
    dal_al, elenco = _anni(periodi) if periodi else ("[DA COMPILARE: periodi d'imposta]", "[DA COMPILARE: periodi d'imposta]")

    # tabella «OGGETTO»
    cella = d.tables[0].rows[0].cells[1].paragraphs[0]
    riscrivi(cella, [(OGGETTO[tipo], "")])

    # blocco del contribuente
    prefisso = (dati.get("forma_prefisso") or "Ditta ind.le").strip()
    den = _o_da_compilare(dati.get("denominazione", ""), "denominazione")
    luogo = _o_da_compilare(dati.get("luogo", ""), "domicilio fiscale e luogo di esercizio")
    pezzi = [(f"{prefisso} ", ""), den, (" con domicilio fiscale e luogo di esercizio in ", ""), luogo]
    att, cod = (dati.get("attivita") or "").strip(), (dati.get("codice_attivita") or "").strip()
    if att:
        pezzi += [(", esercente “", ""), (att, "i"), ("”", "")]
    if cod:
        pezzi += [(f" – cod. attività {cod}", "")]
    pezzi += [(";", "")]
    riscrivi(P[I_DITTA], pezzi)
    riscrivi(P[I_CF], [("Codice Fiscale: ", ""), _o_da_compilare(dati.get("cf", ""), "codice fiscale")])
    if (dati.get("piva") or "").strip():
        riscrivi(P[I_PIVA], [("Partita IVA: ", ""), (dati["piva"].strip(), "")])
    else:
        P[I_PIVA]._p.getparent().remove(P[I_PIVA]._p)
    titolo = (dati.get("titolo_destinatario") or "Sig.").strip()
    dest = _o_da_compilare(dati.get("destinatario", ""), "destinatario")
    riscrivi(P[I_AL], [("AL\t", ""), (f"{titolo}  ", ""), dest])
    riscrivi(P[I_VIA], [("\t", ""), _o_da_compilare(dati.get("indirizzo_destinatario", ""), "indirizzo del destinatario")])

    # formula d'invito
    ora = (dati.get("ora") or "").strip() or ORA_VUOTA
    giorno = (dati.get("data") or "").strip() or DATA_VUOTA
    if tipo == "verifica":
        intro = ("Al fine di consentire a questo Reparto di intraprendere una verifica ai fini delle imposte sui redditi "
                 f"e/o dell'I.V.A. e/o degli altri tributi {NORME_VERIFICA[0].lower() + NORME_VERIFICA[1:]}")
        intro = ("Al fine di consentire a questo Reparto di intraprendere una verifica ai fini delle imposte sui redditi "
                 "e/o dell'I.V.A. e/o degli altri tributi ai sensi dell’art. 32, comma 1, n. 2) e 3) del D.P.R. 29 "
                 "settembre 1973, n. 600 e/o dell’art. 51, comma 2, n. 2) e 3), del D.P.R. 26 ottobre 1972, n. 633 e/o "
                 "dell’art. 2 del D.Lgs 68/2001")
    else:
        intro = ("Al fine di consentire a questo Reparto di intraprendere un controllo fiscale ai fini di P.T., ai sensi e "
                 "per gli effetti degli artt. 52 e 63 del D.P.R. 26 ottobre 1972, n. 633, 33 del D.P.R. 29 settembre "
                 f"1973, n. 600, 2 del D.Lgs 68/2001, nonché della L. n. 4/1929, relativo ai periodi d’imposta {elenco}")
    riscrivi(P[I_FINE], [(f"{intro}, nei confronti del contribuente in oggetto specificato, si invita la S.V. a comparire "
                          f"di persona - o a mezzo rappresentante assistito da procura speciale - alle ore {ora} del "
                          f"giorno {giorno}, presso la sede del Reparto in intestazione.", "")])

    # documenti da recare: elenco numerato (si clona il primo punto)
    docs = _lista(dati.get("documenti")) or ["[DA COMPILARE: documentazione da recare]"]
    primo = P[I_DOC_PRIMO]
    for extra in P[I_DOC_PRIMO + 1:I_DOC_ULTIMO + 1]:
        extra._p.getparent().remove(extra._p)
    riscrivi(primo, [(docs[0], "dc" if docs[0].startswith("[DA COMPILARE") else "")])
    ultimo = primo._p
    for voce in docs[1:]:
        clone = copy.deepcopy(primo._p)
        ultimo.addnext(clone)
        ultimo = clone
        par = docx.text.paragraph.Paragraph(clone, primo._parent)
        riscrivi(par, [(voce, "")])
    riscrivi(P[I_PERIODI], [(f"riferita agli anni d’imposta {dal_al}.", "")])

    # contatti
    ref = _lista(rep.get("referenti"))
    contatto = ("con il " + " o il ".join(ref)) if ref else "con [DA COMPILARE: militari di riferimento]"
    riscrivi(P[I_CONTATTI], [("Nell’ipotesi in cui non fosse possibile ottemperare al presente invito nei modi e nei tempi "
                              "sopra indicati, ovvero nel caso in cui la parte si rende disponibile ad anticipare il giorno "
                              "di presentazione, la S.V. è pregata di prendere sollecitamente contatti ", ""),
                             (contatto, "dc" if not ref else ""), (f" - telefono: {rep['telefono']}.", "")])

    # ragioni giustificative e operazioni
    motivo = (dati.get("motivazione") or "").strip()
    prefisso_r = RAGIONI_VERIFICA if tipo == "verifica" else RAGIONI_CONTROLLO
    if motivo:
        riscrivi(P[I_RAGIONI], [(prefisso_r, ""), (motivo if motivo.endswith(".") else motivo + ".", "")])
    else:
        riscrivi(P[I_RAGIONI], [(prefisso_r, ""), ("[DA COMPILARE: motivo dell'intervento]", "dc")])
    operazione = "verifica" if tipo == "verifica" else "controllo"
    riscrivi(P[I_OPERAZIONI], [(f"2. le operazioni di {operazione} prenderanno in esame i periodi d’imposta {elenco};", "")])

    # firma
    riscrivi(P[I_COMANDANTE], [("IL COMANDANTE DELLA COMPAGNIA" + (" in s.v." if rep.get("in_sv") else ""), "")])
    com = (rep.get("comandante") or "").strip()
    riscrivi(P[I_NOME_COMANDANTE], [(f"({com})", "")] if com else [("([DA COMPILARE: grado nome cognome del Comandante])", "dc")])

    cp = d.core_properties
    cp.author = cp.last_modified_by = cp.title = cp.subject = cp.comments = cp.keywords = ""
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
