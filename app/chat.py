"""Assistente conversazionale per una pratica: raccoglie motivazioni e obiettivi, chiede i documenti necessari e
porta avanti la redazione degli atti (PVOC del primo giorno, giornate successive, PVC) nell'ordine della circolare.

Il modello non modifica nulla da solo: puo' proporre aggiornamenti del fascicolo (blocco <<AZIONI>> in JSON, validato qui)
e proporre la redazione di un atto, che parte solo quando l'operatore preme il pulsante.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile

from . import analisi, calcoli
from . import pvoc as pvoc_mod

MAX_CARATTERI_DOC = 400_000
MAX_DOC_NEL_PROMPT = 6_000          # per documento, gia' riassunto/troncato
MAX_TOTALE_DOC = 30_000

INIZIO, FINE = "<<AZIONI>>", "<<FINE>>"

SYSTEM_CHAT = """Sei l'assistente istruttore di un militare della Guardia di Finanza (Compagnia di Tarquinia, Sezione \
Operativa Volante) che conduce un controllo o una verifica fiscale. Conversi con lui come in una chat, in italiano, \
con tono professionale e sintetico. Il militare e' l'autore degli atti e ne resta responsabile.

Il tuo compito, in ordine:
1. CAPIRE il caso: ragioni che hanno portato ad aprire il controllo/verifica, tipologia, soggetto, periodi d'imposta, \
tributi, obiettivo (a cosa si vuole arrivare). Fai poche domande alla volta (al massimo tre), le piu' utili.
2. CHIEDERE i documenti e le informazioni necessari a svolgere gli accertamenti: fatture (di vendita/acquisto), registri IVA, \
libro giornale e libri contabili, dichiarazioni (Redditi, IVA), F24, estratti conto, contratti, documentazione specifica \
della tipologia (es. bonus edilizi: titoli abilitativi, SAL, asseverazioni, comunicazioni di cessione). Tieni traccia di cio' \
che e' stato chiesto e di cio' che e' stato acquisito (elenco "richieste").
3. RACCOGLIERE i dati che servono per gli atti (vedi "DATI DEL PVOC DEL PRIMO GIORNO" nel contesto): chiedili all'operatore, \
non inventarli. Scrivi in "pvoc_primo" solo cio' che l'operatore ti ha detto o che risulta dai documenti.
4. PROPORRE il percorso: le fasi applicabili nell'ordine della Circolare 1/2018 (l'ordine e' vincolante; una fase non e' \
proponibile se le precedenti non sono concluse o dichiarate non applicabili). Quando hai gli elementi, proponi la redazione \
del PVOC del primo giorno, poi delle giornate successive (una attivita' per giornata) e infine del PVC.
5. SEGNALARE irregolarita' e coerenze che emergono dai documenti, con riferimento alla norma; se emergono indizi di reato \
(es. possibile superamento delle soglie del D.Lgs. 74/2000) ricorda l'art. 220 disp. att. c.p.p. e la comunicazione di notizia \
di reato (art. 347 c.p.p.).

6. DEDURRE al posto dell'operatore: analizza i documenti acquisiti (prospetto e dati tracciati calcolati dal programma, \
estratti dei documenti) e individua le irregolarita' e le eventuali violazioni, fase per fase, nell'ordine della circolare. \
Per ogni violazione che ritieni sussistere proponi un RISCONTRO: fase in cui va constatata, periodo d'imposta, tipo \
(formale | sostanziale | indizio_reato), descrizione fattuale senza cifre, norma violata, e gli ID dei dati/calcoli tracciati \
che ne quantificano l'importo. Proponi un riscontro solo se i dati lo sostengono: se manca un documento o un dato, dillo e \
chiedilo. Le violazioni proposte vengono constatate nei PVOC delle giornate successive solo dopo la conferma dell'operatore. \
Se ti serve un importo che non c'e', proponi un CALCOLO tra voci tracciate (somma, differenza, percentuale): lo esegue il programma.

Regole inderogabili:
- La Circolare 1/2018 e' la fonte di verita' per le fasi e per i contenuti degli atti; gli esempi del Reparto servono per il lessico.
- Non inventare mai fatti, importi, date, orari, protocolli, nominativi o esiti. Gli importi calcolati si producono solo con i \
calcoli tracciati dell'app: non scrivere cifre calcolate a mente.
- Usa i segnaposto tra parentesi quadre (es. [PERSONA_1], [ENTE_1]) esattamente come li ricevi.
- Se un riferimento normativo non e' certo, scrivilo come "da verificare".
- Non dire di avere redatto un atto: la redazione parte quando l'operatore preme il pulsante che tu proponi.

Alla fine di OGNI risposta puoi aggiungere un solo blocco di aggiornamenti, esattamente in questa forma (JSON valido, \
tutte le chiavi facoltative), che l'operatore non vede:
<<AZIONI>>
{"fascicolo": {"motivazione": "...", "obiettivo": "..."},
 "pvoc_primo": {"campo": "valore"},
 "richieste": [{"voce": "Fatture di acquisto 2023-2025", "stato": "richiesto"}],
 "riscontri": [{"fase": "coerenza_interna", "periodo": "2023", "tipo": "sostanziale", "descrizione": "...", "norma": "...", "importi": ["ID1", "ID2"]}],
 "calcoli": [{"tipo": "differenza", "etichetta": "...", "operandi": ["ID1", "ID2"]}],
 "proposte": [{"fase": "avvio", "giornata": "gg/mm/aaaa", "motivo": "perche' ora"}]}
<<FINE>>
- "fascicolo": motivazione e obiettivo riassunti dalle parole dell'operatore (testo formale, senza abbellimenti).
- "pvoc_primo": chiavi ammesse: __CAMPI__.
- "richieste": stato = richiesto | acquisito | non_disponibile. Riporta l'elenco completo aggiornato solo se cambia.
- "riscontri": solo nuovi riscontri (i gia' presenti sono nel contesto); "importi" = ID di dati/calcoli tracciati esistenti.
- "calcoli": tipo = somma | differenza | percentuale (per la percentuale aggiungi "param": "22"); gli operandi sono ID esistenti; \
l'ID del risultato lo assegna il programma e lo trovi nel contesto alla risposta successiva.
- "proposte": fasi (chiavi dell'elenco fasi del contesto) di cui proponi la redazione ora; "giornata" solo se l'operatore l'ha indicata.
"""


def system() -> str:
    return SYSTEM_CHAT.replace("__CAMPI__", ", ".join(pvoc_mod.CAMPI + ("ivi",)))


# ---------------------------------------------------------------- estrazione testo dai documenti
def _txt(el) -> str:
    return (el.text or "").strip() if el is not None else ""


def _fattura_xml(radice) -> tuple[str, list[tuple[str, str]]] | None:
    """Riassume una FatturaPA (una riga per documento). Solo dati necessari ai riscontri."""
    def l(nome, base=None):
        return (base if base is not None else radice).xpath(f".//*[local-name()='{nome}']")
    if not radice.xpath("//*[local-name()='FatturaElettronicaBody']"):
        return None
    righe, nomi = [], []
    cedente = radice.xpath("//*[local-name()='CedentePrestatore']")
    committente = radice.xpath("//*[local-name()='CessionarioCommittente']")

    def anagr(nodo):
        if not nodo:
            return "?"
        n = nodo[0]
        den = _txt(next(iter(l("Denominazione", n)), None))
        if not den:
            den = (_txt(next(iter(l("Nome", n)), None)) + " " + _txt(next(iter(l("Cognome", n)), None))).strip()
        piva = _txt(next(iter(l("IdCodice", n)), None))
        cf = _txt(next(iter(l("CodiceFiscale", n)), None))
        if den:
            nomi.append((den, "ente" if not l("Nome", n) else "persona"))
        return f"{den or '?'} (P.IVA {piva or '-'}, CF {cf or '-'})"
    for body in radice.xpath("//*[local-name()='FatturaElettronicaBody']"):
        gen = l("DatiGeneraliDocumento", body)
        tipo = _txt(next(iter(l("TipoDocumento", body)), None))
        num = _txt(next(iter(l("Numero", body)), None))
        data = _txt(next(iter(l("Data", body)), None))
        tot = _txt(next(iter(l("ImportoTotaleDocumento", body)), None))
        imp = [_txt(x) for x in l("ImponibileImporto", body)]
        iva = [_txt(x) for x in l("Imposta", body)]
        aliq = [_txt(x) for x in l("AliquotaIVA", body)]
        nat = [_txt(x) for x in l("Natura", body)]
        cau = " ".join(_txt(x) for x in l("Causale", body))[:300]
        righe.append(f"Fattura {tipo} n. {num} del {data}; cedente/prestatore {anagr(cedente)}; "
                     f"cessionario/committente {anagr(committente)}; totale {tot or '-'}; imponibile {', '.join(imp) or '-'}; "
                     f"aliquote {', '.join(aliq) or '-'}; imposta {', '.join(iva) or '-'}"
                     + (f"; natura {', '.join(nat)}" if nat else "") + (f"; causale: {cau}" if cau else ""))
        del gen
    return "\n".join(righe), nomi


def fatture_xml(dati: bytes) -> list[dict]:
    """Dati strutturati delle fatture elettroniche (per i riscontri del programma). Lista vuota se non e' una FatturaPA."""
    from lxml import etree
    try:
        radice = etree.fromstring(dati, parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))
    except etree.XMLSyntaxError:
        return []

    def tutti(base, nome):
        return base.xpath(f".//*[local-name()='{nome}']")

    def primo(base, nome):
        r = tutti(base, nome)
        return _txt(r[0]) if r else ""

    def anagr(nodo):
        if not nodo:
            return {"denominazione": "", "piva": "", "cf": ""}
        n = nodo[0]
        den = primo(n, "Denominazione") or (primo(n, "Nome") + " " + primo(n, "Cognome")).strip()
        return {"denominazione": den, "piva": primo(n, "IdCodice"), "cf": primo(n, "CodiceFiscale")}
    ced = anagr(radice.xpath("//*[local-name()='CedentePrestatore']"))
    com = anagr(radice.xpath("//*[local-name()='CessionarioCommittente']"))

    def num(x):
        try:
            return float(x)
        except ValueError:
            return 0.0
    out = []
    for body in radice.xpath("//*[local-name()='FatturaElettronicaBody']"):
        data = primo(body, "Data")
        riep = [{"aliquota": num(primo(r, "AliquotaIVA")), "imponibile": num(primo(r, "ImponibileImporto")),
                 "imposta": num(primo(r, "Imposta")), "natura": primo(r, "Natura")}
                for r in tutti(body, "DatiRiepilogo")]
        out.append({"numero": primo(body, "Numero"), "data": data, "anno": int(data[:4]) if data[:4].isdigit() else 0,
                    "tipo_doc": primo(body, "TipoDocumento"), "cedente": ced, "cessionario": com,
                    "totale": num(primo(body, "ImportoTotaleDocumento")), "bollo": num(primo(body, "ImportoBollo")),
                    "riepilogo": riep})
    return out


def estrai_testo(nome: str, dati: bytes) -> tuple[str, str, list[tuple[str, str]]]:
    """(tipo, testo, nomi da pseudonimizzare [(nome, persona|ente)]). Solleva ValueError se il formato non e' leggibile."""
    n = nome.lower()
    if n.endswith(".xml") or n.endswith(".p7m") and dati[:5] == b"<?xml":
        from lxml import etree
        try:
            radice = etree.fromstring(dati, parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))
        except etree.XMLSyntaxError as e:
            raise ValueError(f"XML non valido: {e}") from e
        riassunto = _fattura_xml(radice)
        if riassunto:
            return "xml_fattura", riassunto[0], riassunto[1]
        return "testo", " ".join(radice.itertext())[:MAX_CARATTERI_DOC], []
    if n.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ValueError("Lettura dei PDF non disponibile su questo server") from e
        r = PdfReader(io.BytesIO(dati))
        t = "\n".join((pg.extract_text() or "") for pg in r.pages)
        if not t.strip():
            raise ValueError("Il PDF non contiene testo selezionabile (e' una scansione): serve l'OCR, non ancora disponibile.")
        return "pdf", t[:MAX_CARATTERI_DOC], []
    if n.endswith(".docx"):
        import docx
        try:
            d = docx.Document(io.BytesIO(dati))
        except (zipfile.BadZipFile, KeyError) as e:
            raise ValueError("File Word non valido") from e
        parti = [p.text for p in d.paragraphs]
        for t in d.tables:
            parti += ["\t".join(c.text for c in r.cells) for r in t.rows]
        return "docx", "\n".join(parti)[:MAX_CARATTERI_DOC], []
    if n.endswith((".csv", ".tsv", ".txt", ".md")):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                testo = dati.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if n.endswith((".csv", ".tsv")):
            righe = list(csv.reader(io.StringIO(testo), delimiter="\t" if n.endswith(".tsv") else (";" if testo.count(";") > testo.count(",") else ",")))
            return "csv", "\n".join(" | ".join(r) for r in righe)[:MAX_CARATTERI_DOC], []
        return "testo", testo[:MAX_CARATTERI_DOC], []
    raise ValueError("Formato non supportato. Carica XML delle fatture, PDF con testo, Word, CSV o testo.")


# ---------------------------------------------------------------- contesto per il modello
def contesto(p_tipo: str, tipologia_nome: str, fasi: list[dict], d: dict, documenti: list[dict], atti: list[dict],
             pvoc_iniziale: dict, prospetto: str = "", voci: str = "") -> str:
    fasc = d.get("fascicolo", {}) or {}
    sog = d.get("soggetto", {}) or {}
    righe = [f"TIPO DI INTERVENTO: {p_tipo}. TIPOLOGIA: {tipologia_nome}.",
             "SOGGETTO: " + "; ".join((sog.get("persone", []) or []) + (sog.get("enti", []) or [])),
             "IDENTIFICATIVI: " + "; ".join((sog.get("codici_fiscali", []) or []) + (sog.get("partite_iva", []) or [])),
             "VERBALIZZANTI: " + "; ".join(d.get("verbalizzanti", []) or []),
             f"MOTIVAZIONE (fin qui): {fasc.get('motivazione') or 'non ancora indicata'}",
             f"OBIETTIVO (fin qui): {fasc.get('obiettivo') or 'non ancora indicato'}",
             "", "FASI NELL'ORDINE DELLA CIRCOLARE (chiave | titolo | stato | atto):"]
    for f in fasi:
        righe.append(f"- {f['chiave']} | {f['titolo']} | {f['stato']}" + (f" | atto {f['atto']}" if f["atto"] else "")
                     + (" | condizionale" if f["condizionale"] else ""))
    righe += ["", "RICHIESTE DI DOCUMENTI (stato):"]
    rich = fasc.get("richieste") or []
    righe += [f"- {r['voce']}: {r['stato']}" for r in rich] or ["- nessuna ancora"]
    righe += ["", "DATI DEL PVOC DEL PRIMO GIORNO (campo: valore | mancante):"]
    for k in pvoc_mod.CAMPI:
        v = pvoc_iniziale.get(k, "")
        righe.append(f"- {k}: {v if v else 'MANCANTE'}")
    righe += ["", "DOCUMENTI ACQUISITI (estratti):"]
    tot = 0
    for doc in documenti:
        estr = doc["testo"][:MAX_DOC_NEL_PROMPT]
        if tot + len(estr) > MAX_TOTALE_DOC:
            righe.append(f"- {doc['nome']}: ({doc['caratteri']} caratteri, non riportato per limiti di spazio)")
            continue
        tot += len(estr)
        righe.append(f"--- {doc['nome']} ({doc['tipo']}, {doc['caratteri']} caratteri) ---\n{estr}")
    if not documenti:
        righe.append("- nessuno")
    if prospetto:
        righe += ["", "PROSPETTO DELLE FATTURE (calcolato dal programma sui file XML caricati):", prospetto]
    righe += ["", "DATI E CALCOLI TRACCIATI DISPONIBILI (id - etichetta: valore):", voci or "- nessuno"]
    righe += ["", "RISCONTRI (id | fase | periodo | tipo | stato | origine): descrizione [norma] [importi]"]
    ris = d.get("riscontri") or []
    righe += [f"- {r['id']} | {r['fase']} | {r['periodo']} | {r['tipo']} | {r['stato']} | {r['origine']}: {r['descrizione']} "
              f"[{r['norma']}] [{', '.join(r['importi'])}]" for r in ris] or ["- nessuno"]
    righe += ["", "ATTI GIA' REDATTI: " + (", ".join(f"{a['tipo']} {a['giornata'] or ''} ({a['fase']})" for a in atti) or "nessuno")]
    return "\n".join(righe)


# ---------------------------------------------------------------- azioni proposte dal modello
_RE_BLOCCO = re.compile(re.escape(INIZIO) + r"(.*?)" + re.escape(FINE), re.S)


def separa_azioni(testo: str) -> tuple[str, dict | None]:
    """(testo da mostrare, azioni). Blocco assente o non valido -> None, mai un errore."""
    m = _RE_BLOCCO.search(testo)
    visibile = _RE_BLOCCO.sub("", testo)
    if INIZIO in visibile:                                        # blocco aperto e mai chiuso
        visibile = visibile.split(INIZIO)[0]
    visibile = visibile.strip()
    if not m:
        return visibile, None
    grezzo = m.group(1).strip()
    grezzo = re.sub(r"^```(?:json)?|```$", "", grezzo).strip()
    try:
        dati = json.loads(grezzo)
    except json.JSONDecodeError:
        return visibile, None
    return visibile, dati if isinstance(dati, dict) else None


def _mappa(o, f):
    if isinstance(o, str):
        return f(o)
    if isinstance(o, list):
        return [_mappa(x, f) for x in o]
    if isinstance(o, dict):
        return {k: _mappa(v, f) for k, v in o.items()}
    return o


def valida_azioni(az: dict | None, restore, fasi_valide: dict[str, str], id_voci: set[str] | frozenset = frozenset()) -> dict:
    """Filtra le azioni: solo chiavi note, tipi corretti; i segnaposto tornano dati reali. `fasi_valide`: chiave -> atto."""
    if not az:
        return {}
    az = _mappa(az, lambda s: restore(s)[0])
    out: dict = {}
    fasc = az.get("fascicolo")
    if isinstance(fasc, dict):
        out["fascicolo"] = {k: str(v).strip() for k, v in fasc.items() if k in ("motivazione", "obiettivo") and str(v).strip()}
    pv = az.get("pvoc_primo")
    if isinstance(pv, dict):
        ok = {}
        for k, v in pv.items():
            if k == "ivi":
                ok[k] = bool(v)
            elif k in pvoc_mod.CAMPI and isinstance(v, (str, int, float)) and str(v).strip():
                ok[k] = str(v).strip()
        out["pvoc_primo"] = ok
    ri = az.get("richieste")
    if isinstance(ri, list):
        out["richieste"] = [{"voce": str(r["voce"]).strip()[:200], "stato": r.get("stato") if r.get("stato") in ("richiesto", "acquisito", "non_disponibile") else "richiesto"}
                            for r in ri if isinstance(r, dict) and str(r.get("voce", "")).strip()][:60]
    rs = az.get("riscontri")
    if isinstance(rs, list):
        out["riscontri"] = []
        for r in rs:
            if not (isinstance(r, dict) and r.get("fase") in fasi_valide and str(r.get("descrizione", "")).strip()):
                continue
            imp = [i for i in (r.get("importi") or []) if isinstance(i, str) and i in id_voci]
            out["riscontri"].append({
                "fase": r["fase"], "periodo": str(r.get("periodo", "") or "").strip()[:40],
                "tipo": r.get("tipo") if r.get("tipo") in ("formale", "sostanziale", "indizio_reato") else "sostanziale",
                "descrizione": str(r["descrizione"]).strip()[:1500], "norma": str(r.get("norma", "") or "da verificare").strip()[:300],
                "importi": imp, "origine": "ai"})
        out["riscontri"] = out["riscontri"][:20]
    ca = az.get("calcoli")
    if isinstance(ca, list):
        out["calcoli"] = [{"tipo": c["tipo"], "etichetta": str(c.get("etichetta", "") or "").strip()[:200],
                           "operandi": [str(x) for x in c.get("operandi", []) if isinstance(x, str)],
                           "param": str(c.get("param", "") or "")}
                          for c in ca if isinstance(c, dict) and c.get("tipo") in ("somma", "differenza", "percentuale")
                          and isinstance(c.get("operandi"), list) and str(c.get("etichetta", "")).strip()][:10]
    pr = az.get("proposte")
    if isinstance(pr, list):
        out["proposte"] = [{"fase": r["fase"], "giornata": str(r.get("giornata", "") or "").strip()[:10],
                            "motivo": str(r.get("motivo", "") or "").strip()[:300], "atto": fasi_valide[r["fase"]]}
                           for r in pr if isinstance(r, dict) and r.get("fase") in fasi_valide and fasi_valide[r["fase"]]][:6]
    return out


def applica_azioni(d: dict, az: dict) -> None:
    fasc = d.setdefault("fascicolo", {})
    if "fascicolo" in az:
        fasc.update(az["fascicolo"])
    if "pvoc_primo" in az:
        fasc.setdefault("pvoc_primo", {}).update(az["pvoc_primo"])
    if "richieste" in az:
        fasc["richieste"] = az["richieste"]
    if "proposte" in az:
        fasc["proposte"] = az["proposte"]
    if az.get("riscontri"):
        nuovi = [{**r, "chiave": analisi.chiave_ai(r)} for r in az["riscontri"]]
        d["riscontri"] = analisi.unisci_riscontri(d.get("riscontri", []), nuovi + [
            r for r in d.get("riscontri", []) if r["chiave"] not in {n["chiave"] for n in nuovi}])
    for c in az.get("calcoli", []):
        cfg = d.setdefault("calcoli", {"dati": [], "calcoli": []})
        nuovo = {"id": f"C{len(cfg['calcoli']) + 1}", **c}
        try:
            calcoli.registro_da_dati({"dati": cfg["dati"], "calcoli": cfg["calcoli"] + [nuovo]})
        except Exception:                                     # operandi inesistenti o operazione non valida: si scarta
            continue
        cfg["calcoli"].append(nuovo)


# ---------------------------------------------------------------- nomi non in anagrafica (controllo in piu' sul testo libero)
_PAROLA = r"[A-ZÀ-Ý][a-zà-ÿ'’]{2,}|[A-ZÀ-Ý]{3,}"
_RE_NOME = re.compile(rf"\b(?:{_PAROLA})(?:\s+(?:{_PAROLA})){{1,3}}\b")
_NON_NOME = {"guardia", "finanza", "compagnia", "tarquinia", "sezione", "operativa", "volante", "agenzia", "entrate",
             "legge", "decreto", "circolare", "codice", "civile", "penale", "procedura", "statuto", "contribuente",
             "reparto", "comando", "garante", "regione", "lazio", "roma", "italia", "repubblica", "imposta", "valore",
             "aggiunto", "redditi", "persone", "fisiche", "societa", "società", "regime", "forfettario", "superbonus",
             "ecobonus", "sismabonus", "bonus", "edilizi", "quadro", "modello", "processo", "verbale", "operazioni",
             "compiute", "verifica", "controllo", "constatazione", "invito", "presentarsi", "ordine", "accesso",
             "direttore", "comandante", "tenente", "maresciallo", "capitano", "brigadiere", "luogotenente", "vice",
             "ordinario", "dpr", "dlgs", "art", "artt", "comma", "iva", "irpef", "irap", "ires", "cnr", "pvoc", "pvv",
             "pvc", "fatture", "registri", "libro", "giornale", "dichiarazione", "dichiarazioni", "cassetto", "fiscale",
             "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre", "ottobre",
             "novembre", "dicembre", "del", "della", "delle", "dei", "degli", "per", "con", "che", "non", "sono", "come"}


def nomi_sospetti(testo_anonimo: str) -> list[str]:
    """Sequenze di 2-4 parole maiuscole non riconosciute (possibili nomi di persone/enti fuori dall'anagrafica)."""
    pulito = re.sub(r"\[[A-Z_]+_\d+\]", " ", testo_anonimo)
    out = []
    for m in _RE_NOME.finditer(pulito):
        parole = m.group(0).split()
        if all(w.lower() not in _NON_NOME for w in parole):
            out.append(m.group(0))
    return sorted(set(out))
