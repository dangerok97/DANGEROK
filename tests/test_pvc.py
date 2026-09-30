import io
import zipfile

import docx

from app import atti_word, pvc

D = {"data": "21 ottobre 2025", "denominazione": "Alfa Srl", "sede": "Roma", "luogo": "Tarquinia", "rappresentante": "Mario Rossi",
     "nascita": "nato a Roma il 01/01/1970", "residenza": "via Inventata 1", "documento": "carta d'identità n. X", "qualita": "legale rappresentante",
     "cf": "00000000000", "piva": "00000000000", "codice_attivita": "00.00.00", "data_inizio": "05.09.2025", "tributo": "I.V.A.",
     "dal": "01/01/2023", "al": "05/09/2025", "direttore": "Ten. Test", "ufficio": "Ufficio dell’Agenzia delle Entrate di Viterbo",
     "documenti_richiesti": "copia delle fatture", "dichiarazione_parte": "Sottoscrivo", "fogli": "7 (sette)", "allegati": "5 (cinque)"}


def test_parte_fissa_identica_agli_esempi():
    t = pvc.costruisci(D, ["Ten. A", "Mar. B"], sezioni={"contabile": "ok"})
    for f in ("articolo 6-bis", "art. 6, comma 1, del D.lgs. n. 218/1997", "articolo 5-quater del D.lgs. n. 218/1997",
              "art. 13 del D.lgs. 18 dicembre 1997, n. 472", "3 VIOLAZIONI FORMALI".replace("3 ", "3.\t"),
              "La trasmissione all’Ufficio Territoriale", "dall’Ufficio dell’Agenzia delle Entrate di Viterbo",
              "Le ragioni che hanno determinato la scelta del contribuente, sono da ricondursi ad una autonoma attività informativa",
              "“””NULLA”””.", "I VERBALIZZANTI    LA PARTE"):
        assert f in t, f
    assert "[DA COMPILARE" in t                              # sezioni non fornite


def test_separa_sezioni():
    s = pvc.separa_sezioni("=== CONTABILE ===\nuno\n=== SOSTANZIALE ===\ndue\n=== FORMALI ===\ntre\n=== SOSTANZIALI ===\nquattro")
    assert s == {"contabile": "uno", "sostanziale": "due", "formali": "tre", "sostanziali": "quattro"}
    assert pvc.separa_sezioni("testo libero") == {}


def test_word_pvc_tabella_e_centrati():
    sez = {"contabile": "ok", "sostanziale": "ok", "formali": "Nessuna.",
           "sostanziali": ">> PERIODO D'IMPOSTA 2024\nA. Violazioni IVA.\n| | Descrizione della violazione constatata | Fonte normativa della violazione\n"
                          "| a. | Violazione X. | Norma violata: art. 1\nL'autore della violazione è Tizio."}
    b = atti_word.crea_atto("PVC", pvc.costruisci(D, ["Ten. A"], sezioni=sez), con_spiegazioni=False, data="21/10/2025", soggetto="Alfa")
    d = docx.Document(io.BytesIO(b))
    assert len(d.tables) == 1 and len(d.tables[0].rows) == 2 and len(d.tables[0].columns) == 3
    assert d.tables[0].rows[0].cells[1].text.startswith("Descrizione della violazione")
    centrati = [p.text for p in d.paragraphs if p.alignment == 1]
    assert "FATTO" in centrati and "PERIODO D'IMPOSTA 2024" in centrati
    sec = d.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0 and round(sec.left_margin.cm, 2) == 2.0
    assert zipfile.ZipFile(io.BytesIO(b)).testzip() is None



def test_varianti_forma_intervento():
    t = pvc.costruisci({**D, "attivita": "convertito", "data_conversione": "26.06.2025", "ricerche": True, "senza_formali": True,
                        "assistenza": "la saltuaria presenza della parte", "trasmissione_pec": True, "doc_siglata": True},
                       ["Ten. A"], impresa=False)
    assert "successivamente convertito il 26.06.2025 in una verifica fiscale" in t
    assert "sono state effettuate ricerche" in t and "acquisito alla verifica" in t
    assert "3.\tVIOLAZIONI SOSTANZIALI." in t and "4.\tSEZIONE CONCLUSIVA." in t and "VIOLAZIONI FORMALI" not in t
    assert "posta elettronica certificata" in t and "siglata dai verbalizzanti e dalla parte" in t
    v = pvc.costruisci({**D, "attivita": "verifica"}, ["Ten. A"])
    assert "una verifica fiscale ai fini dell'I.V.A." in v and "p.v. di verifica" in v and "3.\tVIOLAZIONI FORMALI." in v
    c = pvc.costruisci(D, ["Ten. A"])
    assert "attività di controllo fiscale" in c and "sin dall’inizio del controllo" in c
