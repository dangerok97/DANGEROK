import io
import re
import zipfile

import docx
import pytest
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm

from app import invito_word as iw

DATI = {"forma_prefisso": "Ditta ind.le", "denominazione": "ROSSI MARIO", "luogo": "Tarquinia (VT) via dei Test, nr. 1",
        "attivita": "COMMERCIO AL DETTAGLIO", "codice_attivita": "47.19.90", "cf": "RSSMRA80A01H501U",
        "piva": "01234567897", "titolo_destinatario": "Sig.", "destinatario": "ROSSI MARIO",
        "indirizzo_destinatario": "Via dei Test n. 1 - TARQUINIA", "periodi": ["2021", "2022", "2023"],
        "documenti": ["Fatture di acquisto", "Fatture di vendita"], "motivazione": "emergono scostamenti"}
REP = {"comandante": "Ten. Nome COGNOME", "in_sv": True, "referenti": ["Lgt. Uno UNO", "Mar. Due DUE"],
       "telefono": "0766/856028"}


def gen(tipo="verifica", dati=None, rep=REP):
    return docx.Document(io.BytesIO(iw.crea_invito({**DATI, **(dati or {})}, tipo, rep)))


def _testo(d):
    return "\n".join(p.text for p in d.paragraphs)


def test_pagina_e_margini_come_il_modello():
    s = gen().sections[0]
    assert (round(s.page_width.cm, 1), round(s.page_height.cm, 1)) == (21.0, 29.7)       # A4, non Letter
    assert (round(s.top_margin.cm, 2), round(s.bottom_margin.cm, 2), round(s.left_margin.cm, 2), round(s.right_margin.cm, 2)) == (1.0, 1.25, 2.0, 2.0)
    assert s.different_first_page_header_footer is True
    assert "pagina" in "".join(p.text for p in s.header.paragraphs)                      # numerazione in testata


def test_logo_e_intestazione_fissi():
    d = gen()
    assert d.paragraphs[0]._p.xpath(".//w:drawing")                                      # stemma nel primo paragrafo
    z = zipfile.ZipFile(io.BytesIO(iw.crea_invito(DATI, "verifica", REP)))
    assert "word/media/image1.png" in z.namelist()
    t = _testo(d)
    for riga in ("Guardia di Finanza", "COMPAGNIA TARQUINIA", "Sezione Operativa Volante",
                 "Via delle Fiamme Gialle n. 1- 01016 Tarquinia (VT)", "vt1120000p@pec.gdf.it", "vt112.protocollo@gdf.it"):
        assert riga in t
    p = d.paragraphs
    assert [(r.font.name, r.font.size.pt, r.bold) for r in p[1].runs if r.text] == [("Arial", 16.0, True)]
    assert [(r.font.name, r.font.size.pt, r.bold) for r in p[3].runs if r.text] == [("Arial", 12.0, True)]


def test_carattere_arial_ovunque():
    d = gen()
    nomi = {r.font.name for p in d.paragraphs for r in p.runs if r.text.strip()}
    assert nomi == {"Arial"}                                                              # mai Times New Roman o altro


def test_tabella_oggetto_per_tipo():
    assert "Avvio di una verifica fiscale ai fini delle imposte sui redditi" in gen("verifica").tables[0].rows[0].cells[1].text
    assert "Avvio di un controllo fiscale ai fini di P.T." in gen("controllo").tables[0].rows[0].cells[1].text
    assert gen().tables[0].rows[0].cells[0].text.strip() == "OGGETTO:"


def test_blocco_contribuente_e_destinatario():
    t = _testo(gen())
    assert "Ditta ind.le ROSSI MARIO con domicilio fiscale e luogo di esercizio in Tarquinia (VT) via dei Test, nr. 1, esercente “COMMERCIO AL DETTAGLIO” – cod. attività 47.19.90;" in t
    assert "Codice Fiscale: RSSMRA80A01H501U" in t and "Partita IVA: 01234567897" in t
    assert "AL\tSig.  ROSSI MARIO" in t and "\tVia dei Test n. 1 - TARQUINIA" in t
    corsivo = [r.text for p in gen().paragraphs for r in p.runs if r.italic and r.text.strip()]
    assert corsivo == ["COMMERCIO AL DETTAGLIO"]                                         # l'attivita' in corsivo come nell'originale


def test_senza_partita_iva_e_senza_attivita():
    t = _testo(gen(dati={"piva": "", "attivita": "", "codice_attivita": ""}))
    assert "Partita IVA" not in t and "esercente" not in t and "cod. attività" not in t


def test_formule_fisse_verifica_e_controllo():
    v, c = _testo(gen("verifica")), _testo(gen("controllo"))
    assert "art. 32, comma 1, n. 2) e 3) del D.P.R. 29 settembre 1973, n. 600" in v
    assert "artt. 52 e 63 del D.P.R. 26 ottobre 1972, n. 633" in c and "L. n. 4/1929" in c
    assert "2. le operazioni di verifica prenderanno in esame i periodi d’imposta 2021, 2022, 2023;" in v
    assert "2. le operazioni di controllo prenderanno in esame i periodi d’imposta 2021, 2022, 2023;" in c
    for t in (v, c):                                   # blocchi comuni identici al modello
        assert "riferita agli anni d’imposta dal 2021 al 2023." in t
        assert "▪ \tassistere personalmente alle operazioni di controllo;" in t
        assert "da Euro 258,00 a Euro 2065,00" in t and "RELAZIONE DI NOTIFICAZIONE" in t
        assert "I NOTIFICATORI" in t and "IL NOTIFICATO" in t
        assert "alle ore ... del giorno …………, presso la sede del Reparto in intestazione." in t


def test_ragioni_giustificative():
    assert "in quanto emergono scostamenti." in _testo(gen("verifica"))
    c = _testo(gen("controllo", dati={"motivazione": "sull’ottemperanza dell’inversione contabile"}))
    assert "deve eseguire un controllo sull’ottemperanza dell’inversione contabile." in c


@pytest.mark.parametrize("n", [1, 3, 8])
def test_elenco_documenti_numerato(n):
    docs = [f"Documento {i}" for i in range(1, n + 1)]
    d = gen(dati={"documenti": docs})
    voci = [p for p in d.paragraphs if p.text in docs]
    assert [p.text for p in voci] == docs and all(p._p.xpath(".//w:numPr") for p in voci)   # stessa numerazione del modello


def test_periodi_consecutivi_e_non():
    assert "dal 2019 al 2023" in _testo(gen(dati={"periodi": ["2019", "2020", "2021", "2022", "2023"]}))
    t = _testo(gen(dati={"periodi": ["2019", "2021", "2026 fino al 16/02"]}))
    assert "riferita agli anni d’imposta 2019, 2021, 2026 fino al 16/02." in t


def test_firma_e_contatti_da_impostazioni():
    t = _testo(gen())
    assert "IL COMANDANTE DELLA COMPAGNIA in s.v." in t and "(Ten. Nome COGNOME)" in t
    assert "contatti con il Lgt. Uno UNO o il Mar. Due DUE - telefono: 0766/856028." in t
    senza = _testo(gen(rep={**REP, "in_sv": False}))
    assert "IL COMANDANTE DELLA COMPAGNIA" in senza and "in s.v." not in senza


def test_dati_mancanti_evidenziati_in_turchese_senza_errori():
    d = gen(dati={"denominazione": "", "cf": "", "destinatario": "", "motivazione": ""}, rep={})
    dc = [r.text for p in d.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.TURQUOISE]
    assert any("denominazione" in x for x in dc) and any("codice fiscale" in x for x in dc)
    assert any("destinatario" in x for x in dc) and any("motivo" in x for x in dc)
    assert any("Comandante" in x for x in dc) and any("militari di riferimento" in x for x in dc)


def test_il_modello_non_contiene_dati_personali():
    xml = zipfile.ZipFile(iw.MODELLO).read("word/document.xml").decode()
    for k in ("SANTORO", "SNTGNR", "03497950984", "Licchetta", "RUSSO", "CEFAL", "Pompeo", "292030"):
        assert k not in xml
    assert docx.Document(str(iw.MODELLO)).core_properties.author == ""


def test_nessun_segnaposto_del_modello_nel_documento():
    t = _testo(gen()) + gen().tables[0].rows[0].cells[1].text
    assert "{{" not in t and "}}" not in t and "DOC1" not in t
