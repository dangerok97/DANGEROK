import io

import docx
import pytest

from app import atti_word

TESTO = """PROCESSO VERBALE DI OPERAZIONI COMPIUTE
Il giorno [DA COMPILARE: data] in Tarquinia. {{SPIEGA: art. 52 DPR 633/72}}
VERBALIZZANTI
Ten. [MILITARE_1]
PARTE
- acquisito il documento A

I VERBALIZZANTI    LA PARTE
"""
ATTESI = {"PVOC": (1.27, 1.00, 1.27, 1.27), "PVV": (0.75, 1.00, 2.00, 2.00), "PVC": (2.00, 2.25, 2.00, 2.00),
          "CNR": (2.50, 1.75, 2.00, 2.00)}


def _doc(tipo, **kw):
    return docx.Document(io.BytesIO(atti_word.crea_atto(tipo, TESTO, data="06/03/2026", soggetto="D.I. Prova", **kw)))


@pytest.mark.parametrize("tipo", list(ATTESI))
def test_pagina_font_e_intestazione(tipo):
    d = _doc(tipo)
    s = d.sections[0]
    assert (round(s.page_width.cm, 1), round(s.page_height.cm, 1)) == (21.0, 29.7)
    assert tuple(round(x.cm, 2) for x in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)) == ATTESI[tipo]
    assert s.different_first_page_header_footer
    assert d.styles["Normal"].font.name == "Arial"
    t = [p.text for p in d.paragraphs]
    assert t[1:4] == ["Guardia di Finanza", "COMPAGNIA TARQUINIA", "Sezione Operativa Volante"]
    assert d.paragraphs[0]._p.xpath(".//w:drawing")            # stemma
    corr = s.header.paragraphs[0].text
    assert "06/03/2026" in corr and "D.I. Prova" in corr


def test_titolo_una_volta_e_firme():
    t = [p.text for p in _doc("PVOC").paragraphs]
    assert t.count("PROCESSO VERBALE DI OPERAZIONI COMPIUTE") == 1
    assert "\tI VERBALIZZANTI\tLA PARTE" in t


def test_spiegazioni_gialle_e_versione_pulita():
    con = _doc("PVV")
    gialli = [r.text for p in con.paragraphs for r in p.runs if r.font.highlight_color is not None]
    assert any("art. 52" in g for g in gialli)
    pulito = _doc("PVV", con_spiegazioni=False)
    assert not any("art. 52" in p.text for p in pulito.paragraphs)
    assert any("DA COMPILARE" in p.text for p in pulito.paragraphs)


def test_dati_mancanti_nell_intestazione_corrente():
    d = docx.Document(io.BytesIO(atti_word.crea_atto("PVC", TESTO)))
    assert "DA COMPILARE" in d.sections[0].header.paragraphs[0].text


def test_titolo_e_testo_sulla_stessa_riga_non_si_perde():
    d = docx.Document(io.BytesIO(atti_word.crea_atto("PVOC", "PROCESSO VERBALE DI OPERAZIONI COMPIUTE Il giorno 01/01 in Tarquinia.\nAltro testo.")))
    t = "\n".join(p.text for p in d.paragraphs)
    assert "Il giorno 01/01 in Tarquinia." in t and "Altro testo." in t


def test_markdown_ripulito():
    d = docx.Document(io.BytesIO(atti_word.crea_atto("PVV", "## **FATTO**\nTesto con **grassetto**.")))
    t = "\n".join(p.text for p in d.paragraphs)
    assert "**" not in t and "##" not in t and "FATTO" in t
