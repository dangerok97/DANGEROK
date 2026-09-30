import io

import docx
from docx.enum.text import WD_COLOR_INDEX

from app import wordexport

T = ("PREMESSA\nIl giorno 01/01/2026 si e' presentato [PERSONA_1]. {{SPIEGA: Art. 12 L. 212/2000; dato dagli appunti.}}\n"
     "Ora di chiusura: [DA COMPILARE: orario]. {{SPIEGA: Non presente negli appunti,\nquindi non inventato.}}")


def runs(b):
    d = docx.Document(io.BytesIO(b))
    return [r for p in d.paragraphs for r in p.runs]


def test_spiegazioni_in_giallo_e_dati_mancanti_in_turchese():
    r = runs(wordexport.crea_docx(T, True))
    gialle = [x.text for x in r if x.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert len(gialle) == 2 and all(t.strip().startswith("(Spiegazione:") and t.endswith(")") for t in gialle)
    assert "Art. 12 L. 212/2000" in gialle[0] and "quindi non inventato" in gialle[1]
    assert [x.text for x in r if x.font.highlight_color == WD_COLOR_INDEX.TURQUOISE] == ["[DA COMPILARE: orario]"]
    assert all(x.italic for x in r if x.font.highlight_color == WD_COLOR_INDEX.YELLOW)


def test_versione_pulita_senza_spiegazioni():
    b = wordexport.crea_docx(T, False)
    r = runs(b)
    assert not [x for x in r if x.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    testo = "\n".join(p.text for p in docx.Document(io.BytesIO(b)).paragraphs)
    assert "Spiegazione" not in testo and "Art. 12" not in testo and "[DA COMPILARE: orario]" in testo
    assert "si e' presentato [PERSONA_1]." in testo


def test_pulisci():
    assert wordexport.pulisci("A {{SPIEGA: x}} B") == "A B"


def test_intestazione_in_grassetto_e_nessun_autore():
    b = wordexport.crea_docx(T, True)
    d = docx.Document(io.BytesIO(b))
    assert d.paragraphs[0].runs[0].bold and d.core_properties.author == ""
