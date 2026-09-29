import io
from decimal import Decimal
from types import SimpleNamespace as NS

import docx
from docx.enum.text import WD_COLOR_INDEX
import pytest

from app import ai, calcoli, wordexport
from app.privacy.pseudonymizer import Pseudonymizer


def reg_esempio():
    r = calcoli.Registro()
    r.dato("D1", "Fatturato 2023 (somma fatture n. 1-20)", "96.400,00", "fatture emesse, all. 2")
    r.dato("D2", "Compensi dichiarati 2023", "93.000,00", "quadro LM, rigo LM22, all. 3")
    r.differenza("C1", "Differenza tra fatturato e dichiarato 2023", "D1", "D2")
    return r


def test_formato_e_parsing():
    assert calcoli.euro(Decimal("3400")) == "3.400,00" and calcoli.euro(Decimal("1234567.5")) == "1.234.567,50"
    for s, v in [("3.400,00", "3400.00"), ("€ 1.234,56", "1234.56"), ("3400,5", "3400.5"), ("3400.50", "3400.50")]:
        assert calcoli.parse_importo(s) == Decimal(v)


def test_differenza_come_nell_esempio_dell_utente():
    r = reg_esempio()
    assert r.voci["C1"].valore == Decimal("3400.00")
    sp = r.spiegazione("C1")
    for parte in ("euro 96.400,00", "fatture emesse, all. 2", "euro 93.000,00", "quadro LM, rigo LM22", "= euro 3.400,00"):
        assert parte in sp


def test_percentuali_dell_esempio_forfettario():
    r = calcoli.Registro()
    r.dato("D1", "Compensi non dichiarati", "20,00", "indagini finanziarie")
    r.percentuale("C1", "Base imponibile (coefficiente 78%)", "D1", "78")
    r.percentuale("C2", "Imposta sostitutiva 15%", "C1", "15")
    assert r.voci["C1"].valore == Decimal("15.60") and r.voci["C2"].valore == Decimal("2.34")
    assert "dove C1" in r.spiegazione("C2")            # sotto-calcolo spiegato


def test_arrotondamento_half_up_e_somma():
    r = calcoli.Registro()
    r.dato("D1", "a", "0,005", "x")
    assert r.voci["D1"].valore == Decimal("0.01")
    r.dato("D2", "b", "10,10", "x"); r.dato("D3", "c", "0,20", "x")
    assert r.somma("C1", "tot", ["D2", "D3"]).valore == Decimal("10.30")


def test_operando_inesistente():
    with pytest.raises(calcoli.CalcoloErrore):
        calcoli.Registro().differenza("C1", "x", "D1", "D2")


def test_risolvi_importi_inserisce_valore_e_calcolo_una_volta():
    r = reg_esempio()
    t = calcoli.risolvi_importi("Differenza di {{IMPORTO:C1}}; ripeto {{IMPORTO:C1}}; manca {{IMPORTO:C9}}.", r)
    assert t.count("{{SPIEGA:") == 1 and t.count("euro 3.400,00") == 3 and "C9 non presente" in t


def test_importi_non_tracciati_ignora_orari_e_ammessi():
    r = reg_esempio()
    t = "Ore 10,30. Differenza euro 3.400,00, versati euro 1.234,56, e € 20,00 e euro 7000"
    assert calcoli.importi_non_tracciati(t, r) == ["1.234,56", "20,00", "7.000,00"]
    assert calcoli.importi_non_tracciati(t, r, {"1.234,56", "20,00", "7.000,00"}) == []
    assert calcoli.importi_da_testo("ho versato 1.234,56 euro e euro 500") >= {"1.234,56", "500,00"}


def test_le_spiegazioni_non_contano_come_importi_non_tracciati():
    r = reg_esempio()
    t = calcoli.risolvi_importi("Differenza {{IMPORTO:C1}}", r)
    assert calcoli.importi_non_tracciati(t, r) == []       # nella spiegazione compaiono 96.400,00 e 93.000,00


class FakeAI:
    def __init__(self, testo):
        self.testo = testo
        self.richieste = []
        self.beta = NS(messages=NS(create=self._c))

    def _c(self, **kw):
        self.richieste.append(kw)
        return NS(stop_reason="end_turn", model="m", content=[NS(type="text", text=self.testo)])


def test_ai_usa_solo_importi_tracciati_e_segnala_gli_inventati():
    r = reg_esempio()
    c = FakeAI("Emerge una differenza di {{IMPORTO:C1}} {{SPIEGA: dalla riconciliazione.}} e ulteriori euro 999,00.")
    b = ai.genera_bozza(Pseudonymizer(), istruzione="x", contesto="appunti senza importi", checklist=[],
                        client=c, registro=r)
    prompt = c.richieste[0]["messages"][0]["content"]
    assert "C1 - Differenza tra fatturato e dichiarato 2023: euro 3.400,00" in prompt      # l'AI vede id e valore
    assert "Calcolo C1" in b.testo and "euro 3.400,00" in b.testo                          # il programma inserisce calcolo
    assert b.importi_non_tracciati == ["999,00"] and "DA COMPILARE: importi scritti senza calcolo" in b.testo


def test_importo_scritto_dall_operatore_negli_appunti_e_ammesso():
    c = FakeAI("Sono stati versati euro 500,00.")
    b = ai.genera_bozza(Pseudonymizer(), istruzione="x", contesto="Appunti: versati euro 500,00", checklist=[],
                        client=c, registro=calcoli.Registro())
    assert b.importi_non_tracciati == []


def test_word_riporta_il_calcolo_in_giallo():
    r = reg_esempio()
    t = calcoli.risolvi_importi("Emerge una differenza di {{IMPORTO:C1}} tra fatturato e dichiarato.", r)
    d = docx.Document(io.BytesIO(wordexport.crea_docx(t, True)))
    gialle = [x.text for p in d.paragraphs for x in p.runs if x.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert len(gialle) == 1 and "96.400,00" in gialle[0] and "93.000,00" in gialle[0] and "3.400,00" in gialle[0]
    testo = "\n".join(p.text for p in docx.Document(io.BytesIO(wordexport.crea_docx(t, False))).paragraphs)
    assert testo.strip() == "Emerge una differenza di euro 3.400,00 tra fatturato e dichiarato."


def test_registro_da_dati_e_prompt():
    cfg = {"dati": [{"id": "D1", "etichetta": "A", "valore": "100,00", "fonte": "f"},
                    {"id": "D2", "etichetta": "B", "valore": "40,00", "fonte": "g"}],
           "calcoli": [{"id": "C1", "tipo": "differenza", "etichetta": "A-B", "operandi": ["D1", "D2"], "param": ""}]}
    r = calcoli.registro_da_dati(cfg)
    assert r.voci["C1"].valore == Decimal("60.00") and "C1 - A-B: euro 60,00" in r.elenco_per_prompt()
