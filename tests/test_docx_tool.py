import docx
import pytest

from app.privacy.docx_tool import OggettiIncorporati, anonimizza_docx
from app.privacy.pseudonymizer import Pseudonymizer
from tests.conftest import cf_fittizio, piva_fittizia


def crea(path, testo, con_tabella=True):
    d = docx.Document()
    d.core_properties.author = "Autore Reale"
    d.core_properties.last_modified_by = "Altro Nome"
    d.add_paragraph(testo)
    if con_tabella:
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = f"P.IVA {piva_fittizia()}"
    d.sections[0].header.paragraphs[0].text = "Intestazione: Mario Rossi"
    d.save(path)


def test_anonimizza_corpo_tabella_header_e_metadati(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    crea(src, f"Il Sig. Mario Rossi, C.F. {cf_fittizio()}, scrive a m.rossi@esempio.it")
    p = Pseudonymizer()
    p.aggiungi_persona("Mario Rossi")
    e = anonimizza_docx(str(src), str(dst), p)
    out = docx.Document(str(dst))
    tutto = "\n".join(x.text for x in out.paragraphs) + out.tables[0].cell(0, 0).text + out.sections[0].header.paragraphs[0].text
    for dato in ("Rossi", cf_fittizio(), piva_fittizia(), "esempio.it"):
        assert dato not in tutto
    assert out.core_properties.author == "" and out.core_properties.last_modified_by == ""
    assert e.residui == []


def test_si_ferma_con_immagini(tmp_path):
    from docx.shared import Pt
    import struct, zlib
    png = (b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
           + struct.pack(">I", zlib.crc32(b"IHDR" + struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))))
    idat = zlib.compress(b"\x00\xff\x00\x00")
    png += struct.pack(">I", len(idat)) + b"IDAT" + idat + struct.pack(">I", zlib.crc32(b"IDAT" + idat))
    png += struct.pack(">I", 0) + b"IEND" + struct.pack(">I", zlib.crc32(b"IEND"))
    img = tmp_path / "i.png"
    img.write_bytes(png)
    d = docx.Document()
    d.add_picture(str(img))
    src = tmp_path / "c.docx"
    d.save(src)
    with pytest.raises(OggettiIncorporati):
        anonimizza_docx(str(src), str(tmp_path / "o.docx"), Pseudonymizer())
